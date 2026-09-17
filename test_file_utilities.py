"""Regression tests for portable, non-destructive file utilities."""
import importlib
import subprocess
import sys

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError

from combineMarkdown import combine_markdown_files
from combinePDFS import combine_pdfs
from combineTextFiles import combine_text_files
from docSplice import split_docx
from folderSizer import analyze_folder
from imagekit import make_thumbnail, process_image
from surfaceFiles import process_files

MODULES = ['combineMarkdown', 'combinePDFS', 'combineTextFiles', 'docSplice',
           'surfaceFiles', 'folderSizer', 'epubToMarkdown']


@pytest.mark.parametrize('module', MODULES)
def test_import_has_no_side_effects(module, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    importlib.reload(importlib.import_module(module))
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize('module', MODULES[:-1])
def test_cli_requires_arguments(module, tmp_path):
    result = subprocess.run([sys.executable, '-m', module], cwd=tmp_path,
                            capture_output=True, text=True, check=False)
    assert result.returncode == 2
    assert 'usage:' in result.stderr
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize('combine,suffix', [(combine_markdown_files, '.md'),
                                           (combine_text_files, '.txt')])
def test_combiner_rerun_and_absolute_output(combine, suffix, tmp_path):
    (tmp_path / ('b' + suffix)).write_text('second', encoding='utf-8')
    (tmp_path / ('a' + suffix)).write_text('first', encoding='utf-8')
    output = tmp_path / ('combined' + suffix)
    combine(tmp_path, str(output))
    first = output.read_text()
    combine(tmp_path, str(output))
    assert output.read_text() == first
    assert first.index('first') < first.index('second')
    assert str(tmp_path) not in first


@pytest.mark.parametrize('combine,suffix', [(combine_markdown_files, '.md'),
                                           (combine_text_files, '.txt')])
def test_decode_failure_preserves_previous_output(combine, suffix, tmp_path):
    output = tmp_path / ('combined' + suffix)
    output.write_text('previous')
    (tmp_path / ('bad' + suffix)).write_bytes(b'\xff')
    with pytest.raises(UnicodeDecodeError):
        combine(tmp_path, output)
    assert output.read_text() == 'previous'
    assert not list(tmp_path.glob('.tidydata-*'))


def test_pdf_rerun_preserves_page_count(tmp_path):
    for name, width in [('a.pdf', 100), ('b.PDF', 200)]:
        with PdfWriter() as writer:
            writer.add_blank_page(width=width, height=100)
            writer.write(tmp_path / name)
    output = combine_pdfs(tmp_path)
    combine_pdfs(tmp_path, output)
    assert [int(p.mediabox.width) for p in PdfReader(output).pages] == [100, 200]


def test_bad_pdf_preserves_output(tmp_path):
    output = tmp_path / 'combined.pdf'
    output.write_bytes(b'previous')
    (tmp_path / 'bad.pdf').write_bytes(b'not a PDF')
    with pytest.raises(PdfReadError):
        combine_pdfs(tmp_path)
    assert output.read_bytes() == b'previous'


def test_empty_combiner_and_output_symlink(tmp_path):
    with pytest.raises(ValueError, match='No matching'):
        combine_markdown_files(tmp_path)
    source = tmp_path / 'source.md'
    source.write_text('keep')
    (tmp_path / 'result.md').symlink_to(source)
    with pytest.raises(ValueError, match='symbolic'):
        combine_markdown_files(tmp_path, 'result.md')
    assert source.read_text() == 'keep'


def test_flatten_collisions_existing_files_and_unicode(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    (source / 'a!.txt').write_text('first')
    (source / 'a?.txt').write_text('second')
    nested = source / ('界' * 60)
    nested.mkdir()
    (nested / ('界' * 60 + '.txt')).write_text('unicode')
    destination = tmp_path / 'out'
    destination.mkdir()
    (destination / 'a.txt').write_text('existing')
    copied = process_files(source, destination)
    assert len(copied) == 3
    assert {p.read_text() for p in destination.iterdir()} == {'first', 'second', 'unicode', 'existing'}
    assert all(len(p.name.encode('utf-8')) <= 255 for p in copied)
    process_files(source, destination)
    assert len(list(destination.iterdir())) == 7
    assert (source / 'a!.txt').read_text() == 'first'


def test_flatten_groups_filters_and_symlinks(tmp_path):
    source = tmp_path / 'source'
    nested = source / 'wrapper' / 'Custom category'
    nested.mkdir(parents=True)
    (nested / 'a.pdf').write_bytes(b'pdf')
    (nested / 'a.txt').write_text('text')
    (nested / 'link.pdf').symlink_to(nested / 'a.pdf')
    out = tmp_path / 'out'
    copied = process_files(source, out, extensions=['PDF'], group_depth=1)
    assert len(copied) == 1
    assert copied[0].parent == out / 'Custom category'
    assert copied[0].read_bytes() == b'pdf'


def test_flatten_rejects_overlap_and_destination_group_symlink(tmp_path):
    source = tmp_path / 'source'
    (source / 'category').mkdir(parents=True)
    (source / 'category' / 'a.txt').write_text('text')
    for destination in [source, source / 'out', tmp_path]:
        with pytest.raises(ValueError, match='non-overlapping'):
            process_files(source, destination)
    out = tmp_path / 'out'
    out.mkdir()
    (out / 'category').symlink_to(source / 'category')
    with pytest.raises(ValueError, match='symbolic'):
        process_files(source, out, group_depth=0)


def test_split_markers_and_existing_outputs(tmp_path):
    source = tmp_path / 'source.docx'
    doc = Document()
    for text in ['intro', 'START TWO', 'middle', 'START THREE', 'end']:
        doc.add_paragraph(text)
    doc.save(source)
    outputs = [tmp_path / f'part{i}.docx' for i in range(3)]
    with pytest.raises(ValueError, match='in order'):
        split_docx(source, 'missing', 'START THREE', *outputs)
    assert not any(p.exists() for p in outputs)
    split_docx(source, 'START TWO', 'START THREE', *outputs)
    assert [[p.text for p in Document(f).paragraphs] for f in outputs] == [
        ['intro'], ['START TWO', 'middle'], ['START THREE', 'end']]
    with pytest.raises(ValueError, match='must not already exist'):
        split_docx(source, 'START TWO', 'START THREE', *outputs)


def test_folder_threshold_is_configurable(tmp_path):
    (tmp_path / 'small.txt').write_text('tiny')
    assert analyze_folder(tmp_path)[0] == ''
    assert '.txt' in analyze_folder(tmp_path, min_size=0)[0]


@pytest.mark.parametrize('pipeline', [False, True])
def test_thumbnail_only_strips_source_metadata(tmp_path, pipeline):
    source, target = tmp_path / 'source.png', tmp_path / 'thumb.png'
    exif = Image.Exif()
    exif[315] = 'Private name'
    exif[274] = 6
    Image.new('RGB', (40, 20)).save(source, exif=exif, icc_profile=b'private profile')
    if pipeline:
        process_image(source, thumb_dst=target, thumb_max_width=10)
    else:
        make_thumbnail(source, target, max_width=10)
    with Image.open(target) as img:
        assert img.size == (10, 20)
        assert not img.getexif()
        assert not img.info.get('icc_profile')


def test_thumbnail_extreme_aspect_ratio(tmp_path):
    source, target = tmp_path / 'source.png', tmp_path / 'thumb.png'
    Image.new('RGB', (2000, 1)).save(source)
    assert make_thumbnail(source, target, max_width=1) == (1, 1)


def test_folder_sizer_skips_symlinks(tmp_path):
    source, external = tmp_path / 'source', tmp_path / 'external'
    source.mkdir()
    external.mkdir()
    (external / 'private.txt').write_text('outside')
    (source / 'loop').symlink_to(source, target_is_directory=True)
    (source / 'outside').symlink_to(external, target_is_directory=True)
    (source / 'dangling.txt').symlink_to(tmp_path / 'missing')
    report, _ = analyze_folder(source, min_size=0)
    assert '.txt' not in report
    assert 'outside' not in report
    assert '0.00B' in report


@pytest.mark.parametrize('date_text,expected', [
    ('2024:01:02 03:04:05', '2024-01-02T03:04:05'),
    ('2024:01:02', '2024-01-02T00:00:00'),
    ('invalid date', None),
])
def test_exif_date_preserves_camera_wall_time(tmp_path, date_text, expected):
    from imagekit import extract_exif_date
    source = tmp_path / 'dated.jpg'
    exif = Image.Exif()
    exif[306] = date_text
    Image.new('RGB', (10, 10)).save(source, exif=exif)
    value = extract_exif_date(source)
    assert (value.isoformat() if value else None) == expected
    if value:
        assert value.tzinfo is None
